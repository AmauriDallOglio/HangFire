from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(
    r"C:\Users\OKEA\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-experiment-analysis\assets\reference.docx"
)
OUTPUT = ROOT / "Hangfire_Project_Memo.docx"


def set_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def replace_para(paragraph, text):
    set_text(paragraph, text)
    for run in paragraph.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(10.5)


def add_heading_after(document, anchor_text, heading, paragraphs):
    anchor = None
    for p in document.paragraphs:
        if p.text.strip() == anchor_text:
            anchor = p
            break
    if anchor is None:
        return

    anchor._p.addnext(deepcopy(anchor._p))
    new_heading_p = anchor._p.getnext()
    new_heading = document.paragraphs[[p._p for p in document.paragraphs].index(new_heading_p)]
    new_heading.style = anchor.style
    set_text(new_heading, heading)

    last = new_heading._p
    body_style = document.styles["normal"]
    for text in paragraphs:
        new_p = deepcopy(document.paragraphs[30]._p)
        last.addnext(new_p)
        para = document.paragraphs[[p._p for p in document.paragraphs].index(new_p)]
        para.style = body_style
        set_text(para, text)
        for run in para.runs:
            run.font.name = "Georgia"
            run.font.size = Pt(10.5)
        last = new_p


def clear_table(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                set_text(p, "")


def fill_cell(cell, text, bold=False):
    p = cell.paragraphs[0]
    set_text(p, text)
    for run in p.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(9.5)
        run.bold = bold


def clone_last_row(table):
    row = table.rows[-1]
    tbl = table._tbl
    new_tr = deepcopy(row._tr)
    tbl.append(new_tr)
    return table.rows[-1]


def ensure_rows(table, count):
    while len(table.rows) < count:
        clone_last_row(table)


def write_table(table, rows):
    clear_table(table)
    ensure_rows(table, len(rows))
    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            if c_idx < len(table.rows[r_idx].cells):
                fill_cell(table.rows[r_idx].cells[c_idx], value, bold=(r_idx == 0))


def main():
    doc = Document(str(REFERENCE))

    # Cover/title block.
    for p in doc.paragraphs:
        if p.text == "Experiment Report":
            set_text(p, "Memo do Projeto")
        elif p.text == "Experiment ":
            set_text(p, "Hangfire")
        elif p.text == "Report Name":
            set_text(p, "Memo do Projeto")
        elif p.text == "Author":
            set_text(p, "Preparado para a equipe de desenvolvimento")
        elif p.text == "Month DD, YYYY":
            set_text(p, "16 de setembro de 2026")

    heading_translations = {
        "Document Control": "Controle do Documento",
        "Purpose": "Objetivo",
        "Business Context": "Contexto do Projeto",
        "Experiment Summary": "Resumo da Avaliação",
        "Design Notes": "Funcionamento da Solução",
        "Hypothesis": "Hipótese de Adoção",
        "Success Criteria": "Critérios de Sucesso",
        "Pre-Launch Validation": "Validação Antes do Uso",
        "Sample and Exposure Summary": "Escopo de Uso",
        "Primary Outcome": "Resultado Principal",
        "Guardrail Metrics": "Indicadores de Controle",
        "Segment Results": "Tipos de Jobs",
        "Data Quality and Limitations": "Cuidados e Limitações",
        "Interpretation": "Interpretação",
        "Decision Log": "Registro de Decisões",
        "Post Test Actions": "Próximas Ações",
        "Appendix A: Metric Definitions": "Apêndice A Definições Operacionais",
        "Appendix B: Analyst Notes": "Apêndice B Consultas no SQL Server",
    }
    for p in doc.paragraphs:
        if p.text.strip() in heading_translations:
            set_text(p, heading_translations[p.text.strip()])

    replacements = {
        "Purpose": [
            "Este memo explica o que é o Hangfire, como ele funciona em um projeto .NET e por que ele é útil quando a aplicação precisa executar processamento em segundo plano fora do ciclo de requisição e resposta.",
            "Também descreve a estrutura operacional de um projeto com Hangfire e mostra como localizar as tabelas que ele cria em um banco SQL Server para monitoramento, suporte e investigação de falhas."
        ],
        "Business Context": [
            "Muitas aplicações web começam executando todas as operações de forma síncrona dentro da requisição do usuário. Esse modelo é simples, mas se torna frágil quando o trabalho é demorado, depende de serviços externos ou precisa continuar mesmo depois que o usuário fecha o navegador.",
            "O Hangfire resolve esse problema movendo unidades de trabalho para jobs em segundo plano persistidos. A aplicação enfileira o job rapidamente, devolve o controle ao usuário e deixa workers dedicados executarem a tarefa com retentativas, rastreamento de estado e dashboard de gerenciamento."
        ],
        "Design Notes": [
            "Uma configuração típica do Hangfire tem três partes principais: a API cliente que cria jobs, o armazenamento persistente que guarda dados de jobs e estados, e um ou mais servidores em segundo plano que buscam e executam os jobs nas filas.",
            "Quando o SQL Server é usado como armazenamento, o Hangfire cria um schema, normalmente chamado HangFire, com tabelas para jobs, estados, filas, conjuntos, contadores, hashes, locks e heartbeats dos servidores. Essas tabelas fazem parte do contrato de execução e devem ser inspecionadas por consultas controladas ou pelo dashboard, não editadas manualmente."
        ],
        "Hypothesis": [
            "Se trabalhos recorrentes, agendados e demorados forem movidos para o Hangfire, o projeto deve ganhar execução mais confiável, melhor visibilidade operacional e separação mais limpa entre requisições do usuário e processamento assíncrono."
        ],
        "Success Criteria": [
            "A adoção é bem-sucedida quando tarefas em segundo plano podem ser enfileiradas, executadas, reprocessadas e observadas sem bloquear requisições normais da aplicação. A equipe deve conseguir identificar jobs com falha, ver o histórico de retentativas, localizar as tabelas no SQL Server e confirmar a saúde dos workers pelo dashboard ou por consultas no banco."
        ],
        "Pre-Launch Validation": [
            "Confirmar as versões dos pacotes Hangfire usados pela aplicação e o provedor de armazenamento configurado.",
            "Confirmar que a connection string do SQL Server aponta para o banco correto e que a preparação do schema foi executada com sucesso.",
            "Confirmar que pelo menos um Hangfire Server está em execução no processo da aplicação, worker service, Windows service ou container responsável por processar filas.",
            "Confirmar que o Dashboard está protegido por autorização adequada antes de habilitá-lo fora de um ambiente confiável de desenvolvimento.",
            "Confirmar que jobs de teste executam corretamente nos cenários fire-and-forget, delayed, recurring e continuation."
        ],
        "Sample and Exposure Summary": [
            "O escopo relevante para este projeto é o conjunto de operações que não devem manter uma requisição HTTP aberta. Exemplos comuns incluem envio de emails, geração de relatórios, sincronização de dados, processamento de arquivos enviados, notificação de sistemas externos e rotinas programadas de manutenção."
        ],
        "Primary Outcome": [
            "O principal resultado esperado é confiabilidade operacional. O Hangfire persiste cada job antes da execução, registra transições de estado, faz retentativas em falhas transitórias e expõe o status por meio de dashboard e armazenamento em banco de dados."
        ],
        "Guardrail Metrics": [
            "Os indicadores de controle incluem tamanho das filas, quantidade de jobs com falha, número de retentativas, disponibilidade dos workers, crescimento do SQL Server, controle de acesso ao dashboard e tempo de execução de jobs caros. Esses sinais ajudam a detectar processamento lento, pressão no armazenamento, risco de autorização ou dependências instáveis."
        ],
        "Segment Results": [
            "O Hangfire oferece vários padrões de job. Jobs fire-and-forget executam assim que um worker consegue processá-los. Jobs delayed executam depois de um tempo configurado. Jobs recurring executam em uma agenda. Continuations executam depois que um job pai termina.",
            "O melhor padrão depende da intenção: trabalho acionado pelo usuário costuma usar fire-and-forget, ações futuras planejadas usam delayed, manutenção programada usa recurring e fluxos dependentes usam continuations."
        ],
        "Data Quality and Limitations": [
            "O Hangfire melhora a execução em segundo plano, mas não torna código não idempotente seguro por si só. Jobs devem ser escritos para que retentativas não dupliquem efeitos irreversíveis, como cobrar um cliente ou enviar o mesmo comando externo mais de uma vez.",
            "O armazenamento em SQL Server também exige atenção operacional. Alto volume de jobs pode aumentar o tamanho das tabelas e a carga de escrita, então retenção, índices, desenho de filas e rotinas de limpeza devem ser revisados para produção."
        ],
        "Interpretation": [
            "O Hangfire é uma boa escolha quando o projeto precisa de processamento em segundo plano confiável com baixo custo de implementação. Ele permite usar métodos .NET comuns como jobs, enquanto o framework cuida de persistência, mudanças de estado, polling, retentativas e agendamento.",
            "O principal tradeoff é a responsabilidade operacional. Quando jobs passam a fazer parte do fluxo da aplicação, a equipe precisa monitorar falhas, proteger o dashboard, manter o armazenamento saudável e desenhar handlers com cuidado."
        ],
        "Post Test Actions": [
            "Registrar os serviços do Hangfire e o armazenamento SQL Server na inicialização da aplicação.",
            "Criar filas explícitas para cargas que exigem capacidade ou prioridade diferente de workers.",
            "Adicionar autorização ao dashboard e documentar quem pode acessar dados operacionais de jobs.",
            "Criar um runbook para localizar tabelas do Hangfire no SQL Server e investigar jobs com falha.",
            "Revisar retenção, limpeza, retentativas e políticas de idempotência antes da entrada em produção."
        ],
        "Appendix A: Metric Definitions": [
            "Tamanho da fila: quantidade de jobs aguardando em uma ou mais filas do Hangfire.",
            "Jobs com falha: jobs cujo estado mais recente é Failed e exigem investigação ou nova execução.",
            "Retentativas: novas tentativas de execução após exceções ou falhas transitórias de dependências.",
            "Heartbeat do worker: sinal de que um Hangfire Server está ativo e processando trabalho.",
            "Crescimento do armazenamento: evolução de tamanho e quantidade de linhas nas tabelas Hangfire do SQL Server."
        ],
        "Appendix B: Analyst Notes": [
            "Para localizar as tabelas do Hangfire no SQL Server, primeiro identifique o banco configurado na connection string do provedor Hangfire SQL Server.",
            "No SQL Server Management Studio, expanda o banco, depois Schemas, e procure o schema HangFire. As tabelas mais comuns são HangFire.Job, HangFire.State, HangFire.JobQueue, HangFire.Server, HangFire.Set, HangFire.Hash, HangFire.Counter, HangFire.AggregatedCounter, HangFire.List, HangFire.Lock e HangFire.Schema.",
            "Uma consulta prática de descoberta é: SELECT s.name AS schema_name, t.name AS table_name FROM sys.tables t JOIN sys.schemas s ON t.schema_id = s.schema_id WHERE s.name = 'HangFire' ORDER BY t.name;",
            "Para inspecionar jobs recentes, consulte HangFire.Job e relacione com HangFire.State usando o StateId atual quando apropriado. Para operação de rotina, prefira o Hangfire Dashboard e evite updates ou deletes manuais, salvo por procedimento de manutenção aprovado."
        ],
    }

    current_heading = None
    used = {}
    for p in doc.paragraphs:
        text = p.text.strip()
        replacement_key = text
        if text in heading_translations.values():
            replacement_key = next(k for k, v in heading_translations.items() if v == text)
        if replacement_key in replacements:
            current_heading = replacement_key
            used[current_heading] = 0
            continue
        if current_heading and text.startswith("["):
            idx = used[current_heading]
            values = replacements[current_heading]
            if idx < len(values):
                replace_para(p, values[idx])
            else:
                replace_para(p, "")
            used[current_heading] += 1

    if len(doc.tables) >= 8:
        write_table(doc.tables[0], [
            ["Versão", "Preparado Por", "Revisores", "Data", "Janela de Análise", "Status"],
            ["1.0", "Equipe de desenvolvimento", "Arquitetura e operações", "16 de setembro de 2026", "Planejamento de adoção", "Memo em rascunho"],
        ])
        write_table(doc.tables[1], [
            ["Campo", "Detalhes"],
            ["Nome do Projeto", "Processamento em segundo plano com Hangfire para aplicações .NET"],
            ["Chave do Projeto", "HANGFIRE-SQLSERVER-MEMO"],
            ["Responsável", "Equipe de desenvolvimento da aplicação"],
            ["Objetivo Principal", "Executar e monitorar jobs assíncronos com confiabilidade"],
            ["Armazenamento", "SQL Server com schema HangFire"],
        ])
        write_table(doc.tables[2], [
            ["Componente", "Função", "Observações"],
            ["API cliente", "Enfileira, agenda ou encadeia jobs a partir do código da aplicação.", "Exemplos incluem BackgroundJob.Enqueue, Schedule, ContinueJobWith e RecurringJob.AddOrUpdate."],
            ["Armazenamento", "Persiste jobs, estados, filas, locks, contadores e metadados de servidores.", "O storage SQL Server normalmente cria objetos no schema HangFire."],
            ["Servidor", "Executa workers que consultam filas e processam jobs.", "Pode rodar dentro da aplicação web ou em um processo worker separado."],
            ["Dashboard", "Mostra filas, retentativas, falhas, jobs recorrentes e servidores.", "Deve ser protegido por autorização em ambientes compartilhados ou produtivos."],
        ])
        write_table(doc.tables[3], [
            ["Tipo de Job", "Quando Usar", "Exemplo"],
            ["Fire-and-forget", "Executar trabalho assim que possível após evento de usuário ou sistema.", "Enviar email de confirmação após cadastro."],
            ["Delayed", "Executar trabalho depois de um atraso ou horário específico.", "Enviar lembrete 24 horas após uma ação."],
            ["Recurring", "Executar trabalho em agenda recorrente.", "Gerar relatórios ou executar limpeza à noite."],
            ["Continuation", "Executar trabalho depois que outro job terminar.", "Exportar dados após uma importação bem-sucedida."],
        ])
        write_table(doc.tables[4], [
            ["Área de Validação", "Resultado Esperado"],
            ["Schema SQL Server", "Schema HangFire e tabelas existem no banco configurado."],
            ["Processo worker", "Pelo menos um servidor aparece ativo no dashboard ou na tabela HangFire.Server."],
            ["Processamento de fila", "Jobs de teste passam de queued para succeeded ou failed."],
            ["Retentativas", "Exceções transitórias são reprocessadas conforme a política configurada."],
            ["Segurança", "Acesso ao dashboard fica restrito a usuários autorizados."],
        ])
        write_table(doc.tables[5], [
            ["Objeto SQL Server", "Finalidade"],
            ["HangFire.Job", "Armazena a invocação serializada e os metadados principais do job."],
            ["HangFire.State", "Armazena histórico de estados como Enqueued, Processing, Succeeded, Failed, Scheduled e Deleted."],
            ["HangFire.JobQueue", "Armazena jobs enfileirados aguardando workers."],
            ["HangFire.Server", "Armazena servidores registrados e dados de heartbeat."],
            ["HangFire.Set, Hash, List", "Armazenam metadados de jobs recorrentes, parâmetros e coleções de apoio."],
            ["HangFire.Counter, AggregatedCounter", "Armazenam contadores usados por estatísticas e resumos do dashboard."],
        ])
        write_table(doc.tables[6], [
            ["Decisão", "Justificativa", "Responsável"],
            ["Usar Hangfire para processamento em segundo plano", "Ele oferece jobs persistidos, retentativas, agendamento e dashboard com mudanças moderadas na aplicação.", "Liderança de desenvolvimento"],
            ["Usar storage SQL Server", "A escolha aproveita a operação de banco existente e torna os dados de jobs localizáveis por ferramentas familiares.", "Arquitetura"],
            ["Proteger acesso ao dashboard", "O dashboard expõe dados operacionais e ações como retentativas e exclusões.", "Segurança e operações"],
        ])
        write_table(doc.tables[7], [
            ["Objetivo da Consulta", "SQL"],
            ["Listar tabelas do Hangfire", "SELECT s.name AS schema_name, t.name AS table_name FROM sys.tables t JOIN sys.schemas s ON t.schema_id = s.schema_id WHERE s.name = 'HangFire' ORDER BY t.name;"],
            ["Contar jobs por estado", "SELECT StateName, COUNT(*) AS total FROM HangFire.Job GROUP BY StateName ORDER BY total DESC;"],
            ["Ver servidores ativos", "SELECT Id, Data, LastHeartbeat FROM HangFire.Server ORDER BY LastHeartbeat DESC;"],
        ])

    pbi_steps = [
        "Título da PBI: Implementar processamento assíncrono de Plano de Alarme no Aktian_Backend usando Hangfire para eventos de Solicitação de Serviço e Ordem de Serviço.",
        "Objetivo da PBI: implementar o processo responsável por identificar eventos configurados em Plano de Alarme, localizar planos aplicáveis ao ativo, enviar notificações aos destinatários configurados e registrar o histórico de cada execução na tabela MMS.PlanoAlarmeLog. O processamento deverá ser chamado pelas rotas de SS e OS, porém executado em segundo plano pelo Hangfire após a confirmação da operação principal.",
        "PARTE 1 - RESUMO DA REGRA DE NEGÓCIO NO AKTIAN.",
        "O AKTIAN deverá gerar eventos de Plano de Alarme quando uma Solicitação de Serviço ou Ordem de Serviço for criada ou tiver alteração real de status. A regra principal é simples: se o status anterior for igual ao novo status, nenhum evento deve ser criado e nenhuma notificação deve ser enviada. Se a operação principal sofrer rollback, o evento também não poderá ser processado.",
        "A geração do evento deve ficar centralizada na regra de negócio responsável pela alteração de status. Isso garante que o mesmo comportamento ocorra quando a alteração vier pela API, aplicativo mobile, integração externa ou processamento automático. O desenvolvedor não deve espalhar a regra em vários pontos da aplicação.",
        "Após a gravação bem-sucedida da SS ou OS, o AKTIAN deverá montar um request para /api/v1/PlanoAlarme/RegistrarEvento e enviar esse processamento para o Hangfire. O request deve conter Id_Ativo, PlanoAlarmeTipoCodigo, TipoManutencao, Mensagem, IdRegistroOrigem e NomeTabela.",
        "O PlanoAlarme define o evento monitorado, o tipo de manutenção e o escopo dos ativos. O PlanoAlarmeTipo define o código do evento e quais tipos de manutenção são compatíveis. O Alarme define a mídia. AlarmePlanoAlarme vincula os alarmes ao plano. AlarmeUsuario e AlarmeGrupoUsuario definem os destinatários. As tabelas PlanoAlarmeUnidade, PlanoAlarmeArea, PlanoAlarmeSetor, PlanoAlarmeCentroCusto, PlanoAlarmeLocal, PlanoAlarmeTipoAtivo e PlanoAlarmeAtivo definem os ativos elegíveis.",
        "Para localizar um plano aplicável, o handler deve obter o Ativo pelo Id_Ativo, carregar os vínculos do ativo, localizar planos ativos com PlanoAlarme.Inativo = 0, validar se PlanoAlarmeTipo possui o código recebido em request.PlanoAlarmeTipoCodigo e confirmar se o tipo de manutenção recebido está marcado no plano. Para Corretiva, validar TipoManutencaoCorretiva = 1; para Preditiva, TipoManutencaoPreditiva = 1; para Preventiva, TipoManutencaoPreventiva = 1.",
        "Depois de validar o tipo de evento, o ativo será considerado elegível quando corresponder a uma ou mais configurações de escopo do plano: unidade, área, setor, centro de custo, local, tipo de ativo ou ativo específico. Se nenhum plano aplicável for encontrado, o processo deve encerrar sem registrar erro.",
        "Para cada Plano de Alarme encontrado, o processo deve localizar os alarmes vinculados em AlarmePlanoAlarme, considerar somente alarmes ativos com Alarme.Inativo = 0 e ignorar Alarme.TipoMidia = 0, que representa Desconhecido. Em seguida, deve carregar destinatários vindos de AlarmeUsuario e AlarmeGrupoUsuario, considerando apenas usuários e grupos ativos.",
        "Um usuário repetido na mesma mídia e no mesmo alarme deve receber somente uma notificação para aquele evento. Se um destinatário não possuir o contato exigido pela mídia, como telefone para WhatsApp ou SMS, o envio para esse destinatário deve ser ignorado e registrado, sem interromper os demais envios.",
        "As mídias devem seguir esta regra: 0 Desconhecido não executa envio; 1 Notificação Web armazena a notificação para o usuário visualizar no sistema; 2 Push Mobile envia para dispositivos ativos; 3 E-mail usa o endereço da entidade Usuario; 4 WhatsApp usa o telefone do Usuario; 5 SMS usa o telefone do Usuario.",
        "Para cada tentativa de envio deve ser criado registro em MMS.PlanoAlarmeLog com Id_Tenant, Id_PlanoAlarme, Id_Alarme, RegistroOrigemId, NomeTabela, DataExecucao e Historico. Em caso de sucesso, o histórico deve informar a mídia, o usuário e a mensagem enviada. Em caso de falha, deve informar a mídia, o usuário e o detalhe do erro. Depois de uma execução bem-sucedida, atualizar AlarmePlanoAlarme.DataUltimaExecucao com a data atual.",
        "PARTE 2 - PROCESSO DE INSTALAÇÃO E CLASSES NECESSÁRIAS.",
        "Instalar os pacotes NuGet no projeto Aktian_Backend ou nos projetos da solution responsáveis pela API e aplicação: Hangfire, Hangfire.AspNetCore e Hangfire.SqlServer. Caso a solution use camadas separadas, manter Hangfire.AspNetCore na API e as abstrações necessárias na camada de aplicação/infra, seguindo o padrão já usado no histórico do projeto.",
        "Adicionar a connection string de gravação ou uma connection string específica para o storage do Hangfire no appsettings. O Hangfire criará as tabelas no SQL Server, normalmente no schema HangFire. Não misturar essas tabelas com MMS.PlanoAlarmeLog: HangFire é controle técnico de jobs; MMS.PlanoAlarmeLog é histórico funcional do envio de alarmes.",
        "No Program.cs do Aktian_Backend, registrar builder.Services.AddHangfire(config => config.UseSqlServerStorage(connectionString)); registrar builder.Services.AddHangfireServer(); e mapear o dashboard com app.MapHangfireDashboard ou app.UseHangfireDashboard. Em produção, proteger o dashboard com autorização.",
        "Criar uma classe de configuração, se o padrão do projeto permitir, por exemplo HangfireConfig ou DependencyInjectionHangfire, para concentrar AddHangfire, AddHangfireServer, queues e dashboard. O objetivo é deixar o Program.cs limpo e facilitar manutenção.",
        "Criar o request PlanoAlarmeRegistrarEventoCommandRequest com os campos Id_Ativo, PlanoAlarmeTipoCodigo, TipoManutencao, Mensagem, IdRegistroOrigem e NomeTabela. Criar também PlanoAlarmeRegistrarEventoCommandResponse retornando ResultadoOperacao com mensagem de sucesso ou informação de que nenhum plano aplicável foi encontrado.",
        "Criar o handler PlanoAlarmeRegistrarEventoCommandHandler. Esse handler será a regra principal do processamento: obter ativo, localizar planos ativos, validar tipo de manutenção, carregar alarmes, carregar destinatários, eliminar duplicidades, disparar mídias, registrar PlanoAlarmeLog e atualizar DataUltimaExecucao.",
        "Criar uma classe de serviço para execução via Hangfire, por exemplo PlanoAlarmeHangfireServico. Essa classe deve receber IMediator por injeção de dependência e expor um método como RegistrarEvento(PlanoAlarmeRegistrarEventoCommandRequest request). Dentro desse método, chamar mediator.Send(request).",
        "Criar uma classe ou método centralizador para enfileirar jobs, por exemplo PlanoAlarmeHangfireJob ou IPlanoAlarmeJobService. Essa camada deve encapsular IBackgroundJobClient.Enqueue ou Schedule. As rotas de SS e OS não devem conhecer detalhes internos do Hangfire; elas apenas solicitam o enfileiramento do evento.",
        "Criar interfaces e repositórios necessários para consulta das tabelas: IPlanoAlarmeRepositorio, IAlarmeRepositorio, IAtivoRepositorio, IUsuarioRepositorio quando ainda não existirem, e repositório de PlanoAlarmeLog. Se o Aktian já possuir repositórios equivalentes, reaproveitar os existentes para evitar duplicação.",
        "Criar serviços de envio por mídia. O mínimo esperado é separar a responsabilidade por tipo de envio: notificação web, push mobile, e-mail, WhatsApp e SMS. A regra do Plano de Alarme decide quem deve receber; cada serviço de mídia apenas executa o envio e retorna sucesso ou falha.",
        "PARTE 3 - IMPLEMENTAÇÃO NAS ROTAS E COMO DEVE SER CHAMADO.",
        "Na rota POST /api/v1/SolicitacaoServico/Inserir, após gravar a SS com sucesso e antes de finalizar o fluxo da aplicação, enfileirar o evento de abertura. PlanoAlarmeTipoCodigo = 0. TipoManutencao = Corretiva. NomeTabela = SolicitacaoServico. IdRegistroOrigem = Id da SS criada. Mensagem = 'A Solicitação de Serviço {SSxxxx} foi criada pelo usuário {usuario.Nome}.'.",
        "Na rota /api/v1/SolicitacaoServico/Alterar, carregar o status anterior antes da alteração e comparar com o novo status. Se não houve mudança, não chamar o Plano de Alarme. Se houve mudança e a transação foi confirmada, enfileirar o evento conforme o novo status da SS.",
        "Mapeamento de SS: código 1 para Solicitação de Serviço - Análise; código 2 para Rejeitada; código 3 para Ordem gerada; código 4 para Cancelada; código 5 para Finalizada. Para todos os eventos de SS, TipoManutencao deve ser Corretiva e NomeTabela deve ser SolicitacaoServico.",
        "As mensagens de SS devem seguir o padrão: 'A Solicitação de Serviço {SSxxxx} foi alterada para o status {Status} pelo usuário {usuario.Nome}.'. Para abertura, usar: 'A Solicitação de Serviço {SSxxxx} foi criada pelo usuário {usuario.Nome}.'.",
        "Na rota POST /api/v1/OrdemServico/Inserir, após gravar a OS com sucesso, enfileirar o evento de abertura. PlanoAlarmeTipoCodigo = 6. TipoManutencao deve ser obtido da própria Ordem de Serviço. NomeTabela = OrdemServico. IdRegistroOrigem = Id da OS criada. Mensagem = 'A Ordem de Serviço {OSxxxx} foi criada pelo usuário {usuario.Nome}.'.",
        "Na rota /api/v1/OrdemServico/Alterar, carregar o status anterior antes da alteração e comparar com o novo status. Se o status não mudou, não chamar o Plano de Alarme. Se mudou e a transação foi confirmada, enfileirar o evento de acordo com o novo status da OS.",
        "Mapeamento de OS: código 7 para Agendamento; código 8 para Em Execução; código 9 para Aguardando Insumo; código 10 para Aguardando Inspeção; código 11 para Em Inspeção; código 12 para Inspeção Aprovada; código 13 para Inspeção Reprovada; código 14 para Cancelada; código 15 para Pausada; código 16 para Finalizada; código 17 para Rejeitada.",
        "As mensagens de OS devem seguir o padrão: 'A ordem de serviço {OSxxxx} foi alterada para o status {Status} pelo usuário {usuario.Nome}!'. Para abertura, usar: 'A Ordem de Serviço {OSxxxx} foi criada pelo usuário {usuario.Nome}.'.",
        "A chamada ao Hangfire deve ocorrer somente após a gravação principal ser confirmada. Se o projeto usar Unit of Work ou transação explícita, o enfileiramento deve acontecer depois do commit. Se a transação falhar ou sofrer rollback, nenhum job de Plano de Alarme deve ser criado.",
        "Exemplo de chamada esperada na regra de negócio: montar PlanoAlarmeRegistrarEventoCommandRequest com os dados da SS ou OS e chamar o serviço de enfileiramento. O serviço deve executar BackgroundJob.Enqueue<PlanoAlarmeHangfireServico>(service => service.RegistrarEvento(request));. Caso o projeto use filas, usar uma fila específica como plano-alarme.",
        "A rota /api/v1/PlanoAlarme/RegistrarEvento deve existir como POST autenticado para permitir execução manual ou integração interna quando necessário. Porém, para o fluxo de SS e OS, a preferência é que a regra de negócio chame o comando via mediator/send ou enfileire o job diretamente, evitando dependência HTTP interna desnecessária dentro do mesmo backend.",
        "Critérios de aceite: ao criar SS, deve ser gerado evento código 0; ao alterar status de SS, deve ser gerado o código correspondente de 1 a 5; ao criar OS, deve ser gerado evento código 6; ao alterar status de OS, deve ser gerado o código correspondente de 7 a 17; se o status não mudar, nenhum evento deve ser gerado; se houver rollback, nenhum job deve ser processado; PlanoAlarmeLog deve registrar sucesso e falha por mídia e usuário; uma falha de envio não pode interromper os demais destinatários.",
        "Critérios técnicos de aceite: Hangfire instalado e configurado no Aktian_Backend; dashboard protegido; jobs visíveis no dashboard; tabelas HangFire criadas no SQL Server; handler de Plano de Alarme testado; rotas de SS e OS chamando o enfileiramento nos pontos corretos; e consulta em MMS.PlanoAlarmeLog demonstrando histórico das execuções."
    ]
    add_heading_after(doc, "Próximas Ações", "PBI para Implementação do Plano de Alarme com Hangfire no Aktian_Backend", pbi_steps)

    doc.core_properties.title = "Memo do Projeto Hangfire"
    doc.core_properties.subject = "O que é Hangfire, como funciona, benefícios e localização das tabelas no SQL Server"
    doc.core_properties.author = "OpenAI Codex"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
